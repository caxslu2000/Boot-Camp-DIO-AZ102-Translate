from docx import Document

# 1. Cria um documento Word "em branco" e válido
doc_teste = Document()

# 2. Adiciona alguns parágrafos de teste em inglês (como se fosse uma música)
doc_teste.add_paragraph("I know you're somewhere out there, somewhere far away.")
doc_teste.add_paragraph("I want you back.")
doc_teste.add_paragraph("You are my sunshine, my only sunshine.")

# 3. Salva o arquivo no formato correto que a sua função precisa
doc_teste.save("/content/MUSICA.docx")

print("✅ Arquivo MUSICA.docx real criado com sucesso!")
